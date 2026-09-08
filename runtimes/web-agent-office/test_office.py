import copy
import io
import os
from pathlib import Path
import unittest
import zipfile

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from renderers import render
from schema import InvalidSpec, validate_spec

SLIDES = {"kind": "slides", "title": "产品上线准备", "theme": "blue", "slides": [
    {"layout": "cover", "title": "产品上线准备", "subtitle": "工作台与成果交付"},
    {"layout": "bullets", "title": "首期范围", "bullets": ["保存会话和项目资料", "生成可编辑的 Office 文件", "按任务查看执行记录"]},
    {"layout": "table", "title": "验收安排", "table": {"headers": ["检查项", "验收标准", "状态"], "rows": [
        ["会话", "刷新后保留内容", "待验证"], ["文件", "可打开和下载", "待验证"], ["权限", "隔离不同用户", "待验证"]]}},
    {"layout": "chart", "title": "示例任务分布", "chart": {"type": "bar", "title": "任务数量", "categories": ["文档", "表格", "演示文稿"], "series": [{"name": "任务", "values": [12, 8, 6]}]}}
]}
DOCUMENT = {"kind": "document", "title": "上线检查清单", "sections": [
    {"heading": "核对范围", "paragraphs": ["这份清单用于安排上线前的核对工作。各项检查通过后，再安排小范围验证。"]},
    {"heading": "检查事项", "table": {"headers": ["检查项", "验收标准", "负责人"], "rows": [
        ["会话保存", "刷新后可以继续查看已保存消息", "产品团队"],
        ["文件交付", "文件可以打开，页面内容完整", "开发团队"],
        ["访问权限", "不同用户不能查看彼此的任务", "测试团队"]]}},
    {"heading": "发布安排", "bullets": ["保留上一版的可恢复版本。", "完成小范围验证后再逐步开放。"]}
]}
SHEETS = {"kind": "spreadsheet", "title": "工单处理统计", "sheets": [
    {"name": "汇总", "columns": [{"name": "指标"}, {"name": "数量", "format": "integer"}], "rows": [
        ["收到工单", {"formula": "=SUM(数据!B2:B3)"}], ["已处理", {"formula": "=SUM(数据!C2:C3)"}], ["待处理", {"formula": "=SUM(数据!D2:D3)"}]]},
    {"name": "数据", "columns": [{"name": "月份"}, {"name": "收到", "format": "integer"}, {"name": "已处理", "format": "integer"}, {"name": "待处理", "format": "integer"}],
     "rows": [["九月", 10, 5, {"formula": "=B2-C2"}], ["十月", 12, 6, {"formula": "=B3-C3"}]],
     "chart": {"type": "bar", "title": "每月工单", "category_column": 0, "value_columns": [1, 2]}}
]}


class SchemaTests(unittest.TestCase):
    def test_data_only_contract(self):
        for source in (SLIDES, DOCUMENT, SHEETS):
            self.assertEqual(validate_spec(source), source)
        for mutation in (
            {**DOCUMENT, "code": "print('never run')"},
            {**DOCUMENT, "kind": "shell"},
            {**SLIDES, "slides": [{"layout": "image", "title": "no URL fetch", "url": "http://localhost/"}]},
        ):
            with self.assertRaises(InvalidSpec):
                validate_spec(mutation)

    def test_formula_security(self):
        for formula in ('=WEBSERVICE("https://example.com")', '=HYPERLINK("https://example.com")',
                        "=SUM([book.xlsx]Sheet!A1)", "=cmd|'/C calc'!A1", "=SUM(A1:A99999)"):
            value = copy.deepcopy(SHEETS)
            value["sheets"][0]["rows"][0][1] = {"formula": formula}
            with self.assertRaises(InvalidSpec):
                validate_spec(value)

    def test_bounds_and_names(self):
        value = copy.deepcopy(SHEETS)
        value["sheets"][1]["columns"][1]["name"] = "月份"
        with self.assertRaises(InvalidSpec):
            validate_spec(value)
        value = copy.deepcopy(SLIDES)
        value["slides"][1]["bullets"] = ["long " * 100]
        with self.assertRaises(InvalidSpec):
            validate_spec(value)

    def test_editable_powerpoint(self):
        result = render(SLIDES, preview=False)
        prs = Presentation(io.BytesIO(result["file"]))
        self.assertEqual(len(prs.slides), 4)
        self.assertTrue(any(shape.has_table for shape in prs.slides[2].shapes))
        self.assertTrue(any(shape.has_chart for shape in prs.slides[3].shapes))
        self.assertIn("首期范围", "\n".join(s.text for s in prs.slides[1].shapes if s.has_text_frame))
        with zipfile.ZipFile(io.BytesIO(result["file"])) as archive:
            self.assertFalse(any(name.lower().endswith(".bin") for name in archive.namelist()))

    def test_editable_document(self):
        result = render(DOCUMENT, preview=False)
        doc = Document(io.BytesIO(result["file"]))
        self.assertEqual(doc.paragraphs[0].text, DOCUMENT["title"])
        self.assertEqual(len(doc.tables), 1)
        self.assertEqual(len(doc.tables[0].rows), 4)
        self.assertEqual(doc.paragraphs[0].style.name, "Title")


@unittest.skipUnless(os.environ.get("WEB_AGENT_OFFICE_INTEGRATION") == "1", "enable the packaged Office renderer integration")
class RenderTests(unittest.TestCase):
    def test_preview_and_native_reopen(self):
        out = Path(os.environ["WEB_AGENT_OFFICE_TEST_OUTPUT"])
        out.mkdir(parents=True, exist_ok=True)
        for spec in (SLIDES, DOCUMENT, SHEETS):
            with self.subTest(kind=spec["kind"]):
                result = render(spec)
                (out / ("sample." + result["extension"])).write_bytes(result["file"])
                (out / (spec["kind"] + ".pdf")).write_bytes(result["preview"])
                self.assertTrue(result["preview"].startswith(b"%PDF-"))
                with zipfile.ZipFile(io.BytesIO(result["file"])) as archive:
                    self.assertIsNone(archive.testzip())
                    self.assertFalse(any("vbaProject" in name for name in archive.namelist()))
                if spec["kind"] == "spreadsheet":
                    cached = load_workbook(io.BytesIO(result["file"]), data_only=True)
                    self.assertEqual(cached["汇总"]["B2"].value, 22)
                    self.assertEqual(cached["汇总"]["B3"].value, 11)
                    self.assertEqual(cached["数据"]["D2"].value, 5)
                    cached.close()
                    native = load_workbook(io.BytesIO(result["file"]), data_only=False)
                    self.assertTrue(native["汇总"]["B2"].value.startswith("=SUM("))
                    self.assertEqual(len(native["数据"]._charts), 1)
                    native.close()

    def test_formula_like_text_is_not_executable(self):
        spec = {"kind": "spreadsheet", "title": "文本安全", "sheets": [{
            "name": "文本", "columns": [{"name": "输入"}], "rows": [['=HYPERLINK("https://example.com")'], ["+1"], ["@SUM(A1)"]]}]}
        result = render(spec, preview=False)
        wb = load_workbook(io.BytesIO(result["file"]))
        self.assertEqual(wb.active["A2"].data_type, "s")
        self.assertEqual(wb.active["A2"].value, '=HYPERLINK("https://example.com")')
        wb.close()


if __name__ == "__main__":
    unittest.main()
